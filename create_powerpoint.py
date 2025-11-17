#!/usr/bin/env python3
"""
Convert Legal Fees Research Markdown to PowerPoint Presentation
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
import re

def create_presentation():
    """Create PowerPoint presentation from markdown content"""

    # Read the markdown file
    with open('legal_fees_presentation.md', 'r', encoding='utf-8') as f:
        content = f.read()

    # Create presentation object
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Define colors
    TITLE_COLOR = RGBColor(0, 51, 102)  # Dark blue
    ACCENT_COLOR = RGBColor(0, 112, 192)  # Medium blue
    TEXT_COLOR = RGBColor(51, 51, 51)  # Dark gray

    # Split content into sections
    sections = content.split('\n---\n')

    for section in sections:
        section = section.strip()
        if not section:
            continue

        lines = section.split('\n')

        # Determine slide type based on content
        if section.startswith('# '):
            # Title slide
            create_title_slide(prs, section, TITLE_COLOR, ACCENT_COLOR)
        elif '## Executive Summary' in section or '## Key Findings Summary' in section:
            create_executive_summary_slide(prs, section, TITLE_COLOR, TEXT_COLOR)
        elif '| ' in section and '|---' in section:
            # Table slide
            create_table_slide(prs, section, TITLE_COLOR, TEXT_COLOR)
        elif section.count('**') > 5:
            # Content with highlights
            create_content_slide(prs, section, TITLE_COLOR, TEXT_COLOR, ACCENT_COLOR)
        else:
            # Regular content slide
            create_content_slide(prs, section, TITLE_COLOR, TEXT_COLOR, ACCENT_COLOR)

    # Save the presentation
    prs.save('legal_fees_presentation.pptx')
    print("PowerPoint presentation created successfully: legal_fees_presentation.pptx")

def create_title_slide(prs, content, title_color, accent_color):
    """Create a title slide"""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Extract title and subtitle
    lines = [l.strip() for l in content.split('\n') if l.strip()]
    title_text = lines[0].replace('# ', '')
    subtitle_text = '\n'.join(lines[1:3]) if len(lines) > 1 else ''

    # Add title
    left = Inches(0.5)
    top = Inches(2.5)
    width = Inches(9)
    height = Inches(1.5)

    title_box = slide.shapes.add_textbox(left, top, width, height)
    title_frame = title_box.text_frame
    title_frame.text = title_text

    p = title_frame.paragraphs[0]
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = title_color
    p.alignment = PP_ALIGN.CENTER

    # Add subtitle if present
    if subtitle_text:
        left = Inches(0.5)
        top = Inches(4.2)
        width = Inches(9)
        height = Inches(1)

        subtitle_box = slide.shapes.add_textbox(left, top, width, height)
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = subtitle_text.replace('##', '').replace('**', '')

        p = subtitle_frame.paragraphs[0]
        p.font.size = Pt(24)
        p.font.color.rgb = accent_color
        p.alignment = PP_ALIGN.CENTER

def create_executive_summary_slide(prs, content, title_color, text_color):
    """Create executive summary slide with key points"""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    lines = [l.strip() for l in content.split('\n') if l.strip()]

    # Title
    title_text = None
    for line in lines:
        if line.startswith('##'):
            title_text = line.replace('##', '').strip()
            break

    if title_text:
        add_slide_title(slide, title_text, title_color)

    # Content
    content_lines = []
    for line in lines:
        if line.startswith('###'):
            content_lines.append('\n' + line.replace('###', '').strip())
        elif line.startswith('- ') or line.startswith('* '):
            content_lines.append('  ' + line)
        elif line.startswith('**') and line.endswith('**'):
            content_lines.append('\n' + line)
        elif line and not line.startswith('#'):
            content_lines.append(line)

    if content_lines:
        add_content_box(slide, '\n'.join(content_lines), text_color)

def create_table_slide(prs, content, title_color, text_color):
    """Create a slide with a table"""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    lines = [l.strip() for l in content.split('\n') if l.strip()]

    # Title
    title_text = None
    for line in lines:
        if line.startswith('##'):
            title_text = line.replace('##', '').strip()
            break

    if title_text:
        add_slide_title(slide, title_text, title_color)

    # Extract table
    table_lines = [l for l in lines if '|' in l]
    if len(table_lines) < 3:
        # Not enough for a table, create regular content slide
        create_content_slide(prs, content, title_color, text_color, title_color)
        return

    # Parse table
    headers = [cell.strip() for cell in table_lines[0].split('|') if cell.strip()]
    separator_idx = next((i for i, l in enumerate(table_lines) if '|---' in l or '|-' in l), None)

    if separator_idx is None:
        create_content_slide(prs, content, title_color, text_color, title_color)
        return

    data_lines = table_lines[separator_idx + 1:]

    rows = len(data_lines) + 1  # +1 for header
    cols = len(headers)

    # Add table
    left = Inches(0.5)
    top = Inches(1.8)
    width = Inches(9)
    height = Inches(5)

    table = slide.shapes.add_table(rows, cols, left, top, width, height).table

    # Set column widths
    for i in range(cols):
        table.columns[i].width = Inches(9.0 / cols)

    # Fill headers
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = title_color

        paragraph = cell.text_frame.paragraphs[0]
        paragraph.font.bold = True
        paragraph.font.size = Pt(14)
        paragraph.font.color.rgb = RGBColor(255, 255, 255)
        paragraph.alignment = PP_ALIGN.CENTER

    # Fill data
    for row_idx, line in enumerate(data_lines, start=1):
        cells_data = [cell.strip() for cell in line.split('|') if cell.strip()]
        for col_idx, cell_data in enumerate(cells_data[:cols]):
            if row_idx < rows:
                cell = table.cell(row_idx, col_idx)
                cell.text = cell_data

                paragraph = cell.text_frame.paragraphs[0]
                paragraph.font.size = Pt(12)
                paragraph.font.color.rgb = text_color

def create_content_slide(prs, content, title_color, text_color, accent_color):
    """Create a standard content slide"""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    lines = [l.strip() for l in content.split('\n') if l.strip()]

    # Title
    title_text = None
    for line in lines:
        if line.startswith('##'):
            title_text = line.replace('##', '').strip()
            break

    if title_text:
        add_slide_title(slide, title_text, title_color)

    # Content - collect bullets and text
    content_parts = []
    current_section = None

    for line in lines:
        if line.startswith('###'):
            # Subsection header
            if current_section:
                content_parts.append(current_section)
            current_section = {'header': line.replace('###', '').strip(), 'items': []}
        elif line.startswith('####'):
            # Sub-subsection
            item = '  • ' + line.replace('####', '').strip()
            if current_section:
                current_section['items'].append(item)
            else:
                content_parts.append(item)
        elif line.startswith('- ') or line.startswith('* '):
            item = '• ' + line[2:].strip()
            if current_section:
                current_section['items'].append(item)
            else:
                content_parts.append(item)
        elif line.startswith('**') and line.endswith('**'):
            # Bold header
            text = line.replace('**', '').strip()
            if current_section:
                current_section['items'].append(text)
            else:
                content_parts.append(text)
        elif not line.startswith('#') and not line.startswith('|') and len(line) > 2:
            # Regular text
            if current_section:
                current_section['items'].append(line)
            else:
                content_parts.append(line)

    if current_section:
        content_parts.append(current_section)

    # Format and add content
    formatted_content = []
    for part in content_parts[:15]:  # Limit to avoid overcrowding
        if isinstance(part, dict):
            formatted_content.append('\n' + part['header'])
            for item in part['items'][:5]:  # Limit items per section
                formatted_content.append('  ' + item)
        else:
            formatted_content.append(part)

    if formatted_content:
        add_content_box(slide, '\n'.join(formatted_content), text_color)

def add_slide_title(slide, title_text, color):
    """Add title to slide"""
    left = Inches(0.5)
    top = Inches(0.3)
    width = Inches(9)
    height = Inches(1)

    title_box = slide.shapes.add_textbox(left, top, width, height)
    title_frame = title_box.text_frame
    title_frame.text = title_text[:100]  # Limit title length

    p = title_frame.paragraphs[0]
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.LEFT

def add_content_box(slide, content_text, color):
    """Add content text box to slide"""
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(9)
    height = Inches(5.5)

    content_box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = content_box.text_frame
    text_frame.word_wrap = True

    # Process content with formatting
    paragraphs = content_text.split('\n')

    for i, para_text in enumerate(paragraphs):
        if not para_text.strip():
            continue

        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()

        # Clean up markdown
        clean_text = para_text.replace('**', '').replace('*', '')
        p.text = clean_text

        # Determine font size based on content
        if clean_text.startswith('•'):
            p.font.size = Pt(14)
            p.level = 0
        elif clean_text.startswith('  •') or clean_text.startswith('    •'):
            p.font.size = Pt(12)
            p.level = 1
        else:
            p.font.size = Pt(16)
            if '**' in para_text:
                p.font.bold = True

        p.font.color.rgb = color
        p.space_after = Pt(6)

if __name__ == '__main__':
    create_presentation()
